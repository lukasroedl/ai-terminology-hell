"""Slide "From recognizing content to creating it": the shift from predicting/classifying to
generating (left), and generative AI's branches with text/LLM highlighted (right).

Replaces the original hand-made slide (text-only, off-system fonts, mentioned tokens too early).
Placed after the Generative AI divider; re-running replaces the slide.
"""
import design as d

MARKER = "slide-generative-ai"
ANCHOR = "section-genai"
HEADLINE = "From recognizing content to creating it"   # also the anchor of the Transformer slide

from pptx.dml.color import RGBColor

LLM_CYAN = RGBColor(0x01, 0xD1, 0xCC)   # the deck's cyan, used to pick out "LLM" (user edit)
ICON = "assets/recognise-input.svg"     # stock icon the user placed on the slide

BRANCHES = [("Text  (LLM)", True), ("Images", False), ("Audio", False), ("Video", False)]


def build(prs):
    d.replace_slide(prs, MARKER)
    s = d.content_slide(
        prs, "Part 4 · Generative AI", HEADLINE,
        "The same machinery, put to a new job: not naming what is there, but making something new.",
        index=d.index_after(prs, ANCHOR), marker=MARKER)

    # --- Left: recognise vs. generate -------------------------------------------------------------
    d.card(s, d.LEFT, 1.95, 4.78, 3.0)

    # Row 1: recognise -> a label
    d.label(s, "Recognise, predict or classify", 0.85, 2.07, w=3.0)
    d.rect(s, 0.85, 2.38, 0.62, 0.5, fill=d.TEAL_TINT, line=d.BORDER, radius=0.05)
    d.svg_picture(s, ICON, 0.92, 2.39, 0.49, 0.49)        # icon inserted by the user
    d.line(s, 1.52, 2.63, 1.88, 2.63, color=d.DIM_ON_DARK, width=1, arrow=True)
    d.chip(s, "Model", 1.93, 2.46, 0.85, h=0.34, fill=d.NAVY, color=d.WHITE, size=9.5)
    d.line(s, 2.83, 2.63, 3.19, 2.63, color=d.DIM_ON_DARK, width=1, arrow=True)
    d.chip(s, "cat · 92 %", 3.24, 2.46, 1.05, h=0.34, fill=d.TEAL_TINT, color=d.NAVY, size=9.5)
    d.text(s, "The answer is a label.", 3.23, 2.89, 1.46, 0.17, size=10, color=d.MUTED)

    d.line(s, 0.85, 3.36, 5.15, 3.36, color=d.BORDER, width=0.75)

    # Row 2: generate -> new content
    d.label(s, "Generate", 0.85, 3.5, w=3.0)
    d.chip(s, "“Write a summary …”", 0.85, 3.84, 1.5, h=0.34, fill=d.TEAL_TINT, color=d.NAVY,
           size=9)
    d.line(s, 2.4, 4.01, 2.76, 4.01, color=d.DIM_ON_DARK, width=1, arrow=True)
    d.chip(s, "Model", 2.81, 3.84, 0.85, h=0.34, fill=d.NAVY, color=d.WHITE, size=9.5)
    d.line(s, 3.71, 4.01, 4.07, 4.01, color=d.DIM_ON_DARK, width=1, arrow=True)
    for i, w in enumerate((1.0, 0.85, 1.0, 0.6)):        # freshly written lines
        d.rect(s, 4.12, 3.82 + i * 0.13, w, 0.07, fill=d.MINT if i == 0 else d.BORDER)
    d.text(s, "The answer is new content.", 4.12, 4.35, 1.17, 0.34, size=10, color=d.MUTED)

    # --- Right: the branches of generative AI -----------------------------------------------------
    x = 5.6
    d.card(s, x, 1.95, d.LEFT + d.CONTENT_W - x, 3.0, accent=d.MINT)
    d.label(s, "Different domains of generative AI", x + 0.23, 2.07, w=3.0)

    parent_x, parent_y, parent_w, parent_h = 5.9, 3.1, 1.45, 0.6
    d.rect(s, parent_x, parent_y, parent_w, parent_h, fill=d.NAVY, radius=0.06)
    d.text(s, "Generative\nAI", parent_x, parent_y + 0.09, parent_w, 0.45, font=d.HEADLINE,
           size=12, color=d.WHITE, align=d.PP_ALIGN.CENTER)

    spine, child_x, child_w = 7.6, 7.85, 1.5
    tops = [2.45, 3.05, 3.65, 4.25]
    d.line(s, parent_x + parent_w, parent_y + parent_h / 2, spine, parent_y + parent_h / 2,
           color=d.TEAL, width=1)
    d.line(s, spine, tops[0] + 0.21, spine, tops[-1] + 0.21, color=d.TEAL, width=1)
    for top, (name, highlight) in zip(tops, BRANCHES):
        d.line(s, spine, top + 0.21, child_x, top + 0.21, color=d.TEAL, width=1)
        d.chip(s, name, child_x, top, child_w, h=0.42, size=10,
               fill=d.MINT if highlight else d.WHITE, color=d.NAVY,
               line=None if highlight else d.BORDER)
    d.rect(s, d.LEFT, 5.11, 0.05, 0.24, fill=d.TEAL)
    d.rich_text(s, [("An ", d.NAVY), ("LLM", LLM_CYAN),
                    (" is one kind of generative model. GenAI is characterized by creating content "
                     "of any kind, including images, audio and video.", d.NAVY)],
                0.78, 5.14, 8.90, 0.19, size=11)
    return s


if __name__ == "__main__":
    prs = __import__("pptx").Presentation("ai-terminology-hell.pptx")
    for slide in list(prs.slides):          # drop the original hand-made slide once
        if any(sh.has_text_frame and "Generative AI produces new outputs" in sh.text_frame.text
               for sh in slide.shapes):
            d.delete_slide(prs, slide)
    build(prs)
    prs.save("ai-terminology-hell.pptx")
    print("Generative AI slide rebuilt")
