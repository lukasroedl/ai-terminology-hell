"""Design system for the deck (see CLAUDE.md -> Visual design system).

Shared colors, fonts and building blocks so every slide script produces the same look as slide 1.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# Colors
NAVY = RGBColor(0x0D, 0x1B, 0x2A)
NAVY_LIGHT = RGBColor(0x14, 0x29, 0x3D)
TEAL = RGBColor(0x1B, 0x9A, 0xAA)
MINT = RGBColor(0x06, 0xD6, 0xA0)
LIGHT_BG = RGBColor(0xF4, 0xF7, 0xF8)
BORDER = RGBColor(0xDD, 0xE5, 0xE9)
MUTED = RGBColor(0x6B, 0x7B, 0x8C)
SOFT_ON_DARK = RGBColor(0xAF, 0xC2, 0xCE)
DIM_ON_DARK = RGBColor(0x8F, 0xA6, 0xB2)  # secondary text on navy, still readable over circles
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL_TINT = RGBColor(0xD4, 0xEC, 0xEF)  # light fills for chips/boxes on light slides
MINT_TINT = RGBColor(0xCF, 0xF5, 0xEA)

# Fonts
HEADLINE = "Oswald"
BODY = "Roboto"
MONO = "JetBrains Mono"

DEFAULT_LAYOUT = 0  # "DEFAULT": no placeholders


def new_slide(prs, background=NAVY, index=None):
    """Add a slide with a solid background; optionally move it to position `index` (0-based)."""
    slide = prs.slides.add_slide(prs.slide_layouts[DEFAULT_LAYOUT])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = background
    if index is not None:
        move_slide(prs, slide, index)
    return slide


def move_slide(prs, slide, index):
    ids = prs.slides._sldIdLst
    for sld_id in ids:
        if prs.slides.get(int(sld_id.get("id"))) is slide:
            ids.remove(sld_id)
            ids.insert(index, sld_id)
            return


def delete_slide(prs, slide):
    ids = prs.slides._sldIdLst
    for sld_id in ids:
        if prs.slides.get(int(sld_id.get("id"))) is slide:
            prs.part.drop_rel(sld_id.rId)
            ids.remove(sld_id)
            # Renumber slide parts to slide1..N, otherwise the next added slide reuses a taken
            # part name and the saved file contains duplicate zip entries.
            prs.part.rename_slide_parts([s.rId for s in ids])
            return


def text(slide, value, x, y, w, h, font=BODY, size=16, color=WHITE, bold=False,
         spacing=None, rotation=0, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """Add a text box at (x, y) inches. `spacing` is letter spacing in 1/100 pt (kicker labels use 300)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.rotation = rotation
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        f = run.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        if spacing is not None:
            run._r.get_or_add_rPr().set("spc", str(spacing))
    return box


def kicker(slide, value, x, y, w=6.0, color=MINT):
    """Small uppercase mono label above a headline."""
    return text(slide, value.upper(), x, y, w, 0.3, font=MONO, size=11, color=color, spacing=300)


def accent_bar(slide, x, y, w=0.9, color=MINT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Emu(41148))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def circle(slide, x, y, diameter, color=TEAL, alpha=18, line_color=None, line_width=0.75):
    """Signature translucent circle from slide 1: `alpha`% fill plus a thin solid outline."""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    c.fill.solid()
    c.fill.fore_color.rgb = color
    srgb = c.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(alpha * 1000)}))
    c.line.color.rgb = line_color or color
    c.line.width = Pt(line_width)
    c.shadow.inherit = False
    return c


# --- Content slides -------------------------------------------------------------------------------
# Geometry of the light content slides (measured from the existing deck).
LEFT = 0.62          # left edge of content
CONTENT_W = 8.90     # content width
TEXT_DARK = NAVY_LIGHT


def replace_slide(prs, marker):
    """Delete the slide that carries a shape named `marker` (previous version of a slide script)."""
    for slide in list(prs.slides):
        if any(sh.name == marker for sh in slide.shapes):
            delete_slide(prs, slide)


def index_after(prs, needle):
    """0-based position right after the first slide containing the text or shape name `needle`."""
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.name == needle or (sh.has_text_frame and needle in sh.text_frame.text):
                return i + 1
    raise SystemExit(f"Anchor slide '{needle}' not found")


def content_slide(prs, kicker_text, headline, subtitle=None, index=None, marker=None):
    """Light content slide: navy side bar, teal kicker, Oswald headline, optional Roboto subtitle.
    The headline shape is named `marker` so a slide script can find and replace its slide."""
    s = new_slide(prs, background=LIGHT_BG, index=index)
    rect(s, 0, 0, 0.22, 5.625, fill=NAVY)
    kicker(s, kicker_text, LEFT, 0.36, w=CONTENT_W, color=TEAL)
    h = text(s, headline, LEFT, 0.66, CONTENT_W, 0.72, font=HEADLINE, size=32, color=NAVY)
    if marker:
        h.name = marker
    if subtitle:
        text(s, subtitle, LEFT, 1.42, CONTENT_W, 0.5, size=14, color=TEXT_DARK)
    return s


def rect(slide, x, y, w, h, fill=None, line=None, rounded=False, line_width=0.75, radius=None):
    """Rectangle in inches. `rounded` uses a relative corner (18% of the short side); `radius` sets a
    fixed corner radius in inches instead."""
    rounded = rounded or radius is not None
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = 0.18 if radius is None else min(0.5, radius / min(w, h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    shape.shadow.inherit = False
    return shape


def card(slide, x, y, w, h, accent=TEAL, fill=WHITE):
    """White card with light border and a thin colored accent on the left (as on the deck's cards)."""
    rect(slide, x, y, w, h, fill=fill, line=BORDER)
    if accent is not None:
        rect(slide, x, y, 0.06, h, fill=accent)


def label(slide, value, x, y, w=4.0, color=TEAL, size=9.5):
    """Small uppercase mono label inside a card or diagram."""
    return text(slide, value.upper(), x, y, w, 0.25, font=MONO, size=size, color=color, spacing=100)


def chip(slide, value, x, y, w, h=0.34, fill=TEAL, color=WHITE, font=MONO, size=11, line=None):
    """Rounded box with centered text (tokens, words, tags)."""
    box = rect(slide, x, y, w, h, fill=fill, line=line, rounded=True)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=TEAL, width=1.25, arrow=False, dashed=False):
    """Straight connector in inches; `arrow=True` adds an arrowhead at the end."""
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        from pptx.enum.dml import MSO_LINE
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def takeaway(slide, value, y=5.02):
    """Bottom takeaway line with a teal marker bar (as on the deck's content slides)."""
    rect(slide, LEFT, y + 0.02, 0.05, 0.24, fill=TEAL)
    text(slide, value, LEFT + 0.16, y, CONTENT_W - 0.16, 0.3, size=12, color=NAVY)


def run(build, default_deck="ai-terminology-hell.pptx"):
    """Standard entry point for slide scripts: build into the deck given on the command line."""
    import sys
    from pptx import Presentation
    deck = sys.argv[1] if len(sys.argv) > 1 else default_deck
    prs = Presentation(deck)
    slide = build(prs)
    prs.save(deck)
    print(f"Wrote slide {list(prs.slides).index(slide) + 1} to {deck}")


def decision(slide, keyword, condition, x, y, w, h=0.44, keyword_size=9, condition_size=10.5):
    """Compact flowchart decision: rounded box with teal outline, mono keyword (IF / ELSE IF) + condition."""
    box = rect(slide, x, y, w, h, fill=WHITE, line=TEAL, rounded=True, line_width=1.25)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for value, font, color, bold in ((keyword + "  ", MONO, TEAL, True), (condition, BODY, NAVY, False)):
        run = p.add_run()
        run.text = value
        run.font.name = font
        run.font.size = Pt(keyword_size if font == MONO else condition_size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def step_label(slide, n, value, x, y, w=4.0):
    """Numbered step heading in the deck's pattern "STEP 1: TEXT"."""
    return label(slide, f"Step {n}: {value}", x, y, w=w, size=8.5)


def step_box(slide, x, y, w, h):
    """Light boundary that groups one step inside a card."""
    return rect(slide, x, y, w, h, fill=LIGHT_BG, line=BORDER, radius=0.08)

