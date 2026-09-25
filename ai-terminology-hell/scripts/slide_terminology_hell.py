"""Slide 2 "Terminology hell": a deliberately chaotic cloud of every term the deck explains.

Re-running replaces the existing slide instead of adding a duplicate.
"""
import sys

from pptx import Presentation

import design as d

DECK = sys.argv[1] if len(sys.argv) > 1 else d.DECK
MARKER = "terminology-hell-title"  # shape name used to find this slide again
POSITION = 1  # directly after the title slide

# Large circle the cloud lives in: center (x, y) and radius, in inches. Fully on the slide.
CIRCLE_X, CIRCLE_Y, CIRCLE_R = 7.25, 2.80, 2.70
MARGIN = 0.12  # minimum distance (in) between a term and the circle edge

# Approximate text width per character per pt, measured from rendered slides (fonts aren't
# installed locally, so exact metrics aren't available).
CHAR_WIDTH = {d.HEADLINE: 0.40, d.BODY: 0.47, d.MONO: 0.55}

# Each entry: (term, font, size pt, color, x in, y in, rotation deg).
# Every term here must be explained somewhere in the deck.
# Big = the core abbreviations of the story; small/dim = supporting terms.
TERMS = [
    ("LLM",                  d.HEADLINE, 40, d.TEAL,         6.00, 0.60, -4),
    ("Token",                d.MONO,     12, d.DIM_ON_DARK,  6.95, 0.88,  0),
    ("MCP",                  d.HEADLINE, 34, d.WHITE,        7.70, 0.58,  4),
    ("Prompt",               d.MONO,     13, d.MINT,         5.00, 1.55, -8),
    ("AI",                   d.HEADLINE, 60, d.WHITE,        5.95, 1.15,  0),
    ("Deep Learning",        d.HEADLINE, 22, d.SOFT_ON_DARK, 7.20, 1.25,  0),
    ("Neural Network",       d.BODY,     16, d.SOFT_ON_DARK, 7.00, 1.85,  0),
    ("GPT",                  d.HEADLINE, 30, d.TEAL,         8.65, 1.70,  3),
    ("Pre-training",         d.MONO,     11, d.MINT,         4.85, 2.05,  0),
    ("Embedding",            d.MONO,     11, d.DIM_ON_DARK,  4.80, 2.45,  0),
    ("Machine Learning",     d.HEADLINE, 20, d.WHITE,        6.20, 2.35,  0),
    ("Transformer",          d.HEADLINE, 20, d.SOFT_ON_DARK, 8.10, 2.38,  0),
    ("Next-word prediction", d.BODY,     13, d.SOFT_ON_DARK, 4.85, 2.95,  0),
    ("GenAI",                d.HEADLINE, 32, d.TEAL,         6.75, 2.85,  0),
    ("RAG",                  d.HEADLINE, 36, d.MINT,         8.20, 2.85,  5),
    ("Tool use",             d.HEADLINE, 20, d.WHITE,        5.10, 3.55,  0),
    ("AI Agent",             d.HEADLINE, 22, d.TEAL,         6.30, 3.55,  0),
    ("MCP Server",           d.BODY,     13, d.SOFT_ON_DARK, 7.60, 3.65, -3),
    ("Hallucination",        d.BODY,     15, d.MINT,         5.50, 4.20, -4),
    ("Agentic AI",           d.HEADLINE, 28, d.WHITE,        6.85, 4.10, -3),
    ("Context window",       d.MONO,     12, d.SOFT_ON_DARK, 5.90, 4.75,  0),
    ("Rule-based AI",        d.MONO,     11, d.DIM_ON_DARK,  7.30, 4.72,  0),
]


def term_box(term, font, size):
    """Width and height (in) of a term's text box."""
    return len(term) * size * CHAR_WIDTH[font] / 72, size / 72 * 1.4


def check_inside_circle():
    """Fail loudly if any term's (rotated) box pokes out of the large circle."""
    import math
    for term, font, size, _, x, y, rot in TERMS:
        w, h = term_box(term, font, size)
        cx, cy, a = x + w / 2, y + h / 2, math.radians(rot)
        for dx, dy in ((-w / 2, -h / 2), (w / 2, -h / 2), (-w / 2, h / 2), (w / 2, h / 2)):
            px = cx + dx * math.cos(a) - dy * math.sin(a)
            py = cy + dx * math.sin(a) + dy * math.cos(a)
            if math.hypot(px - CIRCLE_X, py - CIRCLE_Y) > CIRCLE_R - MARGIN:
                raise SystemExit(f"'{term}' sticks out of the circle; adjust its position")

def build(prs):
    # Remove a previous version of this slide.
    for slide in list(prs.slides):
        if any(sh.name == MARKER for sh in slide.shapes):
            d.delete_slide(prs, slide)

    s = d.new_slide(prs, background=d.NAVY, index=POSITION)

    # Signature circles: the large teal one holds the cloud, the mint one bleeds off the corner.
    d.circle(s, CIRCLE_X - CIRCLE_R, CIRCLE_Y - CIRCLE_R, 2 * CIRCLE_R, color=d.TEAL, alpha=18)
    d.circle(s, 8.3, 3.9, 3.0, color=d.MINT, alpha=12)

    # Left column: kicker -> headline -> body -> accent bar (same rhythm as slide 1).
    d.kicker(s, "Before we start", 0.8, 1.35, w=3.4)
    title = d.text(s, "Terminology\nhell", 0.8, 1.7, 3.4, 1.8, font=d.HEADLINE, size=48)
    title.name = MARKER
    d.text(s, "All of these terms show up in this talk. By the end, each one will have its place "
              "in the story.", 0.8, 3.55, 3.3, 0.9, size=14, color=d.SOFT_ON_DARK)
    d.accent_bar(s, 0.8, 4.6)

    # Right: the term cloud. Boxes are sized roughly to the text (so rotation pivots around the
    # term itself) and don't wrap, so a term never breaks across lines.
    for term, font, size, color, x, y, rot in TERMS:
        w, h = term_box(term, font, size)
        d.text(s, term, x, y, w, h, font=font, size=size, color=color, rotation=rot, wrap=False)
    return s


if __name__ == "__main__":
    check_inside_circle()
    prs = Presentation(DECK)
    build(prs)
    prs.save(DECK)
    print(f"Terminology hell slide written to {DECK} at position {POSITION + 1}")
