"""Design system for the deck (see CLAUDE.md -> Visual design system).

Shared colors, fonts and building blocks so every slide script produces the same look as slide 1.
"""
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# Where the deck lives: the project folder (this file is in its scripts/ subfolder), so scripts
# work no matter which folder they are started from.
PROJECT = Path(__file__).resolve().parent.parent
DECK = str(PROJECT / "ai-terminology-hell.pptx")


def asset(path):
    """A path inside the project folder (e.g. "assets/avatar-ai.svg"), resolved absolutely."""
    return str(PROJECT / path)


# Colors
NAVY = RGBColor(0x0D, 0x1B, 0x2A)
NAVY_LIGHT = RGBColor(0x14, 0x29, 0x3D)
TEAL = RGBColor(0x1B, 0x9A, 0xAA)
MINT = RGBColor(0x06, 0xD6, 0xA0)
LIGHT_BG = RGBColor(0xF4, 0xF7, 0xF8)
BORDER = RGBColor(0xDD, 0xE5, 0xE9)
MUTED = RGBColor(0x6B, 0x7B, 0x8C)
SOFT_ON_DARK = RGBColor(0xAF, 0xC2, 0xCE)
DIM_ON_DARK = RGBColor(0x8F, 0xA6, 0xB2)  # secondary text on navy, still readable over circles
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICON_TEAL = RGBColor(0x1B, 0x9A, 0xAA)   # matches the stock AI icon
TEAL_TINT = RGBColor(0xD4, 0xEC, 0xEF)  # light fills for chips/boxes on light slides
MINT_TINT = RGBColor(0xCF, 0xF5, 0xEA)

# Fonts
HEADLINE = "Oswald"
BODY = "Roboto"
MONO = "JetBrains Mono"

DEFAULT_LAYOUT = 0  # "DEFAULT": no placeholders


def new_slide(prs, background=NAVY, index=None):
    """Add a slide with a solid background; optionally move it to position `index` (0-based)."""
    slide = prs.slides.add_slide(prs.slide_layouts[DEFAULT_LAYOUT])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = background
    if index is not None:
        move_slide(prs, slide, index)
    return slide


def move_slide(prs, slide, index):
    ids = prs.slides._sldIdLst
    for sld_id in ids:
        if prs.slides.get(int(sld_id.get("id"))) is slide:
            ids.remove(sld_id)
            ids.insert(index, sld_id)
            return


def delete_slide(prs, slide):
    ids = prs.slides._sldIdLst
    for sld_id in ids:
        if prs.slides.get(int(sld_id.get("id"))) is slide:
            prs.part.drop_rel(sld_id.rId)
            ids.remove(sld_id)
            # Renumber slide parts to slide1..N, otherwise the next added slide reuses a taken
            # part name and the saved file contains duplicate zip entries.
            prs.part.rename_slide_parts([s.rId for s in ids])
            return


def text(slide, value, x, y, w, h, font=BODY, size=16, color=WHITE, bold=False,
         spacing=None, rotation=0, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """Add a text box at (x, y) inches. `spacing` is letter spacing in 1/100 pt (kicker labels use 300)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.rotation = rotation
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        f = run.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        if spacing is not None:
            run._r.get_or_add_rPr().set("spc", str(spacing))
    return box


def kicker(slide, value, x, y, w=6.0, color=MINT):
    """Small uppercase mono label above a headline."""
    return text(slide, value.upper(), x, y, w, 0.3, font=MONO, size=11, color=color, spacing=300)


def accent_bar(slide, x, y, w=0.9, color=MINT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Emu(41148))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def circle(slide, x, y, diameter, color=TEAL, alpha=18, line_color=None, line_width=0.75):
    """Signature translucent circle from slide 1: `alpha`% fill plus a thin solid outline."""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    c.fill.solid()
    c.fill.fore_color.rgb = color
    srgb = c.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(alpha * 1000)}))
    c.line.color.rgb = line_color or color
    c.line.width = Pt(line_width)
    c.shadow.inherit = False
    return c


# --- Content slides -------------------------------------------------------------------------------
# Geometry of the light content slides (measured from the existing deck).
LEFT = 0.62          # left edge of content
CONTENT_W = 8.90     # content width
TEXT_DARK = NAVY_LIGHT


def replace_slide(prs, marker):
    """Delete the slide that carries a shape named `marker` (previous version of a slide script)."""
    for slide in list(prs.slides):
        if any(sh.name == marker for sh in slide.shapes):
            delete_slide(prs, slide)


def index_after(prs, needle):
    """0-based position right after the first slide containing the text or shape name `needle`."""
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.name == needle or (sh.has_text_frame and needle in sh.text_frame.text):
                return i + 1
    raise SystemExit(f"Anchor slide '{needle}' not found")


def content_slide(prs, kicker_text, headline, subtitle=None, index=None, marker=None,
                  subtitle_size=14):
    """Light content slide: navy side bar, teal kicker, Oswald headline, optional Roboto subtitle.
    The headline shape is named `marker` so a slide script can find and replace its slide."""
    s = new_slide(prs, background=LIGHT_BG, index=index)
    rect(s, 0, 0, 0.22, 5.625, fill=NAVY)
    kicker(s, kicker_text, LEFT, 0.36, w=CONTENT_W, color=TEAL)
    h = text(s, headline, LEFT, 0.66, CONTENT_W, 0.72, font=HEADLINE, size=32, color=NAVY)
    if marker:
        h.name = marker
    if subtitle:
        text(s, subtitle, LEFT, 1.42, CONTENT_W, 0.5, size=subtitle_size, color=TEXT_DARK)
    return s


def rect(slide, x, y, w, h, fill=None, line=None, rounded=False, line_width=0.75, radius=None,
         rotation=0):
    """Rectangle in inches. `rounded` uses a relative corner (18% of the short side); `radius` sets a
    fixed corner radius in inches instead."""
    rounded = rounded or radius is not None
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = 0.18 if radius is None else min(0.5, radius / min(w, h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    shape.shadow.inherit = False
    shape.rotation = rotation
    return shape


def card(slide, x, y, w, h, accent=TEAL, fill=WHITE):
    """White card with light border and a thin colored accent on the left (as on the deck's cards)."""
    rect(slide, x, y, w, h, fill=fill, line=BORDER)
    if accent is not None:
        rect(slide, x, y, 0.06, h, fill=accent)


def label(slide, value, x, y, w=4.0, color=TEAL, size=9.5):
    """Small uppercase mono label inside a card or diagram."""
    return text(slide, value.upper(), x, y, w, 0.25, font=MONO, size=size, color=color, spacing=100)


def chip(slide, value, x, y, w, h=0.34, fill=TEAL, color=WHITE, font=MONO, size=11, line=None):
    """Rounded box with centered text (tokens, words, tags)."""
    box = rect(slide, x, y, w, h, fill=fill, line=line, rounded=True)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=TEAL, width=1.25, arrow=False, dashed=False):
    """Straight connector in inches; `arrow=True` adds an arrowhead at the end."""
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        from pptx.enum.dml import MSO_LINE
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def takeaway(slide, value, y=5.02, size=12, h=0.3):
    """Bottom takeaway line with a teal marker bar (as on the deck's content slides)."""
    rect(slide, LEFT, y + 0.02, 0.05, h - 0.06, fill=TEAL)
    text(slide, value, LEFT + 0.16, y, CONTENT_W - 0.16, h, size=size, color=NAVY)


def run(build, default_deck=DECK):
    """Standard entry point for slide scripts: build into the deck given on the command line."""
    import sys
    from pptx import Presentation
    deck = sys.argv[1] if len(sys.argv) > 1 else default_deck
    prs = Presentation(deck)
    slide = build(prs)
    prs.save(deck)
    print(f"Wrote slide {list(prs.slides).index(slide) + 1} to {deck}")


def decision(slide, keyword, condition, x, y, w, h=0.44, keyword_size=9, condition_size=10.5):
    """Compact flowchart decision: rounded box with teal outline, mono keyword (IF / ELSE IF) + condition."""
    box = rect(slide, x, y, w, h, fill=WHITE, line=TEAL, rounded=True, line_width=1.25)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for value, font, color, bold in ((keyword + "  ", MONO, TEAL, True), (condition, BODY, NAVY, False)):
        run = p.add_run()
        run.text = value
        run.font.name = font
        run.font.size = Pt(keyword_size if font == MONO else condition_size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def step_label(slide, n, value, x, y, w=4.0):
    """Numbered step heading in the deck's pattern "STEP 1: TEXT"."""
    return label(slide, f"Step {n}: {value}", x, y, w=w, size=8.5)


def step_box(slide, x, y, w, h):
    """Light boundary that groups one step inside a card."""
    return rect(slide, x, y, w, h, fill=LIGHT_BG, line=BORDER, radius=0.08)



def rich_text(slide, parts, x, y, w, h, font=BODY, size=12, align=PP_ALIGN.LEFT, spacing=None):
    """One line of text made of (value, colour) parts, e.g. to colour single numbers in a formula."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for value, color in parts:
        run = p.add_run()
        run.text = value
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if spacing is not None:
            run._r.get_or_add_rPr().set("spc", str(spacing))
    return box


def svg_picture(slide, svg_path, x, y, w, h):
    """Insert an SVG image (e.g. a PowerPoint stock icon) at the given position, in inches.

    python-pptx cannot add SVGs through add_picture, so the image part and the <p:pic> element are
    built by hand, the same way PowerPoint writes them (an svgBlip with no raster fallback)."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml import parse_xml

    package = slide.part.package
    used = [p.partname for p in package.iter_parts() if str(p.partname).startswith("/ppt/media/")]
    n = 1
    while PackURI(f"/ppt/media/svg{n}.svg") in used:
        n += 1
    part = Part(PackURI(f"/ppt/media/svg{n}.svg"), "image/svg+xml",
                package=package, blob=open(asset(svg_path), "rb").read())
    rId = slide.part.relate_to(part, RT.IMAGE)

    shape_id = max([sh.shape_id for sh in slide.shapes] or [1]) + 1
    xml = (
        '<p:pic xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<p:nvPicPr><p:cNvPr id="{shape_id}" name="Picture {shape_id}"/>'
        '<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
        '<p:blipFill><a:blip><a:extLst>'
        '<a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
        f'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="{rId}"/>'
        '</a:ext></a:extLst></a:blip><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{Inches(x)}" y="{Inches(y)}"/>'
        f'<a:ext cx="{Inches(w)}" cy="{Inches(h)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
    slide.shapes._spTree.append(parse_xml(xml))


def send_to_back(slide, *shapes):
    """Move shapes to the back of the drawing order (as PowerPoint's "send to back" does)."""
    tree = slide.shapes._spTree
    for i, shape in enumerate(shapes):
        el = shape._element
        tree.remove(el)
        tree.insert(2 + i, el)          # after nvGrpSpPr and grpSpPr


def delete_slides_with_text(prs, needle):
    """Delete every slide that contains `needle` in any text (drops an original hand-made slide)."""
    for slide in list(prs.slides):
        if any(sh.has_text_frame and needle in sh.text_frame.text for sh in slide.shapes):
            delete_slide(prs, slide)


def link(shape, url):
    """Turn all runs of a text shape into a hyperlink to `url`. The link keeps the run's own colour
    (PowerPoint's hlinkClr extension) instead of the theme's blue hyperlink colour."""
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.hyperlink.address = url
            click = r._r.rPr.find(qn("a:hlinkClick"))
            ext_lst = click.makeelement(qn("a:extLst"), {})
            ext = ext_lst.makeelement(qn("a:ext"), {"uri": "{A12FA001-AC4F-418D-AE19-62706E023703}"})
            ext.append(click.makeelement(
                "{http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor}hlinkClr",
                {"val": "tx"}))
            ext_lst.append(ext)
            click.append(ext_lst)
